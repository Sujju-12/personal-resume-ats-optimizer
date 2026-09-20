from io import BytesIO
import fitz
from docx import Document

def parse_resume(filename: str, content: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        doc = fitz.open(stream=content, filetype="pdf")
        text = "\n".join(page.get_text("text") for page in doc).strip()
        if text:
            return text
    if name.endswith(".docx"):
        doc = Document(BytesIO(content))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts).strip()
        if text:
            return text
    raise ValueError("Supported formats are PDF and DOCX with readable text.")
